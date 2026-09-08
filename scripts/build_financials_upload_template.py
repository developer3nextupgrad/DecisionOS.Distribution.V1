"""Build DecisionOS_Financials_Upload_Template.xlsx for Terry / SCBS."""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill("solid", fgColor="1B4F72")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
TITLE_FONT = Font(bold=True, name="Calibri", size=14, color="1B4F72")
LABEL_FONT = Font(bold=True, name="Calibri", size=11)
BODY_FONT = Font(name="Calibri", size=11)
HINT_FILL = PatternFill("solid", fgColor="F4F6F7")
SAMPLE_FILL = PatternFill("solid", fgColor="EAF2F8")
THIN = Border(
    left=Side(style="thin", color="D5D8DC"),
    right=Side(style="thin", color="D5D8DC"),
    top=Side(style="thin", color="D5D8DC"),
    bottom=Side(style="thin", color="D5D8DC"),
)

MAY, JUN, JUL = date(2026, 5, 31), date(2026, 6, 30), date(2026, 7, 31)

# SCBS Decision-OS reformatted totals (May–Jul 2026)
PL = [
    (MAY, 1_385_414.47, 1_227_471.02, 157_943.45, 0.1140, -15_548.72, -0.0112),
    (JUN, 1_348_262.64, 1_124_965.44, 223_297.20, 0.1656, 31_021.55, 0.0230),
    (JUL, 1_221_734.13, 1_060_362.24, 161_371.89, 0.1321, -43_920.38, -0.0360),
]
BS = [
    (MAY, 231_739.63, 1_866_988.97, 2_211_692.43, 913_235.18),
    (JUN, 332_751.35, 1_899_932.35, 2_222_161.05, 1_526_407.45),
    (JUL, 299_299.27, 1_894_808.31, 2_626_674.11, 2_236_232.08),
]
CF = [
    (MAY, -15_548.72, 274_528.27, 231_739.63, -42_788.64),
    (JUN, 31_021.55, 231_739.63, 332_751.35, 101_011.72),
    (JUL, -43_920.38, 332_751.35, 299_299.27, -33_452.08),
]


def style_header(ws: Worksheet, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32


def autosize(ws: Worksheet, min_w=14, max_w=28) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = 0
        for cell in col:
            if cell.value is None:
                continue
            longest = max(longest, min(len(str(cell.value)), max_w))
        ws.column_dimensions[letter].width = max(min_w, longest + 2)


def write_sheet(ws: Worksheet, headers: list[str], rows: list[list], money_cols: set[int], pct_cols: set[int]) -> None:
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h)
    for r, row in enumerate(rows, 2):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.fill = SAMPLE_FILL
            if c in money_cols and not isinstance(val, date):
                cell.number_format = "#,##0.00"
            if c in pct_cols:
                cell.number_format = "0.00%"
            if isinstance(val, date):
                cell.number_format = "YYYY-MM-DD"
    style_header(ws, len(headers))
    autosize(ws)


def add_readme(ws: Worksheet) -> None:
    rows = [
        ("Topic", "Guidance"),
        ("Purpose", "Fill this workbook and upload it in Decision OS (Simplified import). Sample May–Jul 2026 rows are from Steve’s / SCBS reformatted files — replace with your live numbers."),
        ("Required layout", "Row 1 = column headers. Each later row = one month. Date is in the first column (Week_End_Date). Do not put account names down the left side (native QuickBooks layout)."),
        ("Multiple months", "Keep the same file. Add a new row for the new month (for example 2026-08-31). Re-upload the whole file. Prior months are replaced, not doubled."),
        ("P&L / Balance sheet / Cash flow", "Use the three statement tabs, or put the key totals on Weekly_Financials (recommended). Same Week_End_Date on every tab."),
        ("One file, several tabs", "Put P&L, balance sheet, cash flow, and AR in this single .xlsx. Do not make a separate tab per month."),
        ("Native QuickBooks export", "Will not import. Use this header-on-top layout only."),
        ("AR aging", "QuickBooks GL cannot age invoices. Put open invoices on Accounts_Receivable (customer, open amount, due date or aging bucket). The AR dollar total on the balance sheet is not aging."),
        ("Where to upload", "Decision OS → Operations → Uploads → New Batch → Simplified — single workbook → select tenant → Cadence = Monthly → Continue → choose this .xlsx → Analyze workbook → Verify → Import all periods."),
        ("Classic alternative", "Operations → Uploads → New Batch → Classic → one week at a time → upload each CSV/Excel and set Report type (Financial statement, Accounts receivable, etc.)."),
        ("Do not change", "Keep header names in row 1 exactly as shown (Week_End_Date, Net_Sales, COGS, …). Extra expense columns can be omitted; totals are enough."),
        ("Percent columns", "Use a decimal (0.11 = 11%) or 11 — both are accepted. Do not paste dollar amounts into percent columns."),
        ("Sales & inventory", "For a full 7-KPI dashboard you still need sales and inventory detail (see the full DecisionOS_Simplified_Workbook_Template.xlsx). This file is the financial + AR starter promised on the call."),
    ]
    for r, (a, b) in enumerate(rows, 1):
        ws.cell(r, 1, a).font = LABEL_FONT if r > 1 else HEADER_FONT
        ws.cell(r, 2, b).font = BODY_FONT if r > 1 else HEADER_FONT
        ws.cell(r, 1).fill = HEADER_FILL if r == 1 else HINT_FILL
        ws.cell(r, 2).fill = HEADER_FILL if r == 1 else PatternFill()
        ws.cell(r, 1).font = HEADER_FONT if r == 1 else LABEL_FONT
        ws.cell(r, 2).font = HEADER_FONT if r == 1 else BODY_FONT
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 48 if r > 1 else 22
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def main() -> None:
    wb = Workbook()

    readme = wb.active
    readme.title = "README_Import_Map"
    add_readme(readme)

    wf = wb.create_sheet("Weekly_Financials")
    write_sheet(
        wf,
        [
            "Week_End_Date", "Net_Sales", "COGS", "Gross_Profit", "Gross_Margin_%",
            "Net_Income", "Net_Profit_%", "Inventory_Value_End", "AR_Ending", "AP_Ending",
            "Cash_Ending", "AR_Over_60_%", "AP_Past_Due_%", "Fill_Rate_%", "Notes",
        ],
        [
            [MAY, 1_385_414.47, 1_227_471.02, 157_943.45, 0.1140, -15_548.72, -0.0112, 2_211_692.43, 1_866_988.97, 913_235.18, 231_739.63, None, None, None, "Sample — replace AR_Over_60_% from aged AR"],
            [JUN, 1_348_262.64, 1_124_965.44, 223_297.20, 0.1656, 31_021.55, 0.0230, 2_222_161.05, 1_899_932.35, 1_526_407.45, 332_751.35, None, None, None, "Sample"],
            [JUL, 1_221_734.13, 1_060_362.24, 161_371.89, 0.1321, -43_920.38, -0.0360, 2_626_674.11, 1_894_808.31, 2_236_232.08, 299_299.27, None, None, None, "Add Aug 31 as the next row"],
        ],
        money_cols={2, 3, 4, 6, 8, 9, 10, 11},
        pct_cols={5, 7, 12, 13, 14},
    )

    pl = wb.create_sheet("P_and_L")
    write_sheet(
        pl,
        ["Week_End_Date", "Net_Sales", "COGS", "Gross_Profit", "Gross_Margin_%", "Net_Income", "Net_Profit_%"],
        [[d, ns, cogs, gp, gm, ni, np] for d, ns, cogs, gp, gm, ni, np in PL],
        money_cols={2, 3, 4, 6},
        pct_cols={5, 7},
    )

    bs = wb.create_sheet("Weekly_Balance_Sheet")
    write_sheet(
        bs,
        ["Week_End_Date", "Cash_Ending", "AR_Ending", "Inventory_Value_End", "AP_Ending"],
        [[d, cash, ar, inv, ap] for d, cash, ar, inv, ap in BS],
        money_cols={2, 3, 4, 5},
        pct_cols=set(),
    )

    cf = wb.create_sheet("Weekly_Cash_Flow")
    write_sheet(
        cf,
        ["Week_End_Date", "Net_Income", "Cash_Beginning", "Cash_Ending", "Net_Cash_Change"],
        [[d, ni, beg, end, ch] for d, ni, beg, end, ch in CF],
        money_cols={2, 3, 4, 5},
        pct_cols=set(),
    )

    ar = wb.create_sheet("Accounts_Receivable")
    write_sheet(
        ar,
        [
            "Invoice_ID", "Customer_ID", "Customer_Name", "Invoice_Date", "Due_Date",
            "Original_Amount", "Open_Amount", "Days_Past_Due", "Aging_Bucket", "Collection_Status",
        ],
        [
            ["INV-1001", "CUST-001", "Sample Customer A", date(2026, 6, 1), date(2026, 7, 1), 12_000, 12_000, 30, "1-30", "Open"],
            ["INV-1002", "CUST-002", "Sample Customer B", date(2026, 5, 1), date(2026, 5, 31), 8_500, 8_500, 61, "61-90", "Open"],
            ["INV-1003", "CUST-003", "Sample Customer C", date(2026, 3, 15), date(2026, 4, 14), 4_200, 4_200, 108, "90+", "Escalated"],
        ],
        money_cols={6, 7},
        pct_cols=set(),
    )

    root = Path(__file__).resolve().parents[1]
    targets = [
        root / "src" / "DecisionOS.Distribution.Web" / "wwwroot" / "downloads" / "DecisionOS_Financials_Upload_Template.xlsx",
        root / "_client_docs" / "DecisionOS_Financials_Upload_Template.xlsx",
    ]
    for path in targets:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
